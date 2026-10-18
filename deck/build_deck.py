"""Build the presentation deck from docs/deliverables/deck.md on the official template.

The slide text, tables, speaker notes and timings come from ``docs/deliverables/deck.md``; every measured
number and hash is a ``{{placeholder}}`` resolved from an evidence folder through the registry
``docs/deliverables/placeholders.md`` and the resolver of ``scripts/render_deliverables.py`` (imported, not
copied). Design constants (tier caps, score bands) are read from ``config/`` or typed as text, as in deck.md.

Usage (from the repository root)::

    uv run --with python-pptx --with matplotlib --with qrcode python deck/build_deck.py
    uv run --with python-pptx --with matplotlib --with qrcode python deck/build_deck.py --evidence evidence
    uv run --with python-pptx --with matplotlib --with qrcode python deck/build_deck.py \
        --evidence runs/stale-evidence-20261012 --pdf --previews

``--evidence`` is a folder laid out like ``evidence/``: ``<network>/evidence.json`` (and the rest of an
``acb evidence export``) plus ``attack_suite.json``, ``backtest.json``, ``benign_suite.json`` and
``dialogue_eval.json`` at its top. A network folder itself (one holding ``evidence.json``) also works; the
suite files are then looked up beside it and in its parent. Default: ``evidence/`` once
``evidence/base_sepolia/evidence.json`` exists, otherwise the local Anvil set in
``runs/stale-evidence-20261012``.

A value whose evidence file is missing never prints: the tile, table row, bullet clause or sentence that
uses it is left out, and the value is listed on the console. On Base Sepolia evidence only the optional
evaluations (``dialogue_eval.json``, ``benign_suite.json``) may be missing; any other missing value stops
the build (``--draft`` allows it). On any other network the deck is a draft: each slide carries a DRAFT
line. Before saving, every slide and note is checked for leftovers ("n/a", "uncommitted", an unrendered
``{{placeholder}}`` or a bracketed presenter prompt); any leftover stops the build.
"""

from __future__ import annotations

import argparse
import importlib.util
import json
import re
import shutil
import subprocess
import sys
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import yaml
from lxml import etree
from PIL import ImageFont
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

REPO = Path(__file__).resolve().parent.parent
DECK_DIR = REPO / "deck"
ASSETS = DECK_DIR / "assets"
PREVIEWS = DECK_DIR / "previews"
TEMPLATE = REPO / "private" / "templates" / "presentation-template.pptx"
SOURCE = REPO / "docs" / "deliverables" / "deck.md"
OUT = DECK_DIR / "InternationalAI-TeamMongolia-Deck.pptx"
STALE = REPO / "runs" / "stale-evidence-20261012"
REPO_URL = "https://github.com/taivnaaiwnl/shenzhen-comp"

sys.path.insert(0, str(REPO / "src"))
from acb import reason_codes

_spec = importlib.util.spec_from_file_location("render_deliverables", REPO / "scripts" / "render_deliverables.py")
assert _spec and _spec.loader
rd = importlib.util.module_from_spec(_spec)
sys.modules["render_deliverables"] = rd
_spec.loader.exec_module(rd)

# ---------------------------------------------------------------------------------------------------
# look: template background art, Arial text, Okabe-Ito colours for data (colour-blind safe)
# ---------------------------------------------------------------------------------------------------

FONT = "Arial"
MONO = "Courier New"


def rgb(hexstr: str) -> RGBColor:
    return RGBColor.from_string(hexstr)


INK = "1E2235"
MUTED = "5B6075"
ACCENT = "6779C6"  # the template's number colour
ACCENT_DARK = "3F4C9E"
CARD = "F2F4FB"
CARD_LINE = "CDD3EC"
WHITE = "FFFFFF"
OI_BLUE = "0072B2"
OI_ORANGE = "E69F00"
OI_SKY = "56B4E9"
OI_GREEN = "009E73"
OI_VERM = "D55E00"
OI_YELLOW = "F0E442"
DRAFT_ON_WHITE = "A04500"

EMU_IN = 914400


def In(v: float) -> Emu:
    return Emu(round(v * EMU_IN))


# ---------------------------------------------------------------------------------------------------
# evidence: the renderer's Context with the evidence folder remapped
# ---------------------------------------------------------------------------------------------------


class EvidenceContext(rd.Context):
    """``rd.Context`` reading ``evidence/...`` from another folder (same registry, same extraction rules)."""

    def __init__(self, network: str, groups: dict[str, list[str]], root: Path) -> None:
        super().__init__(network, groups)
        root = root.resolve()
        if (root / "evidence.json").is_file():
            self.net_dir, self.roots = root, [root, root.parent]
        else:
            self.net_dir, self.roots = root / network, [root]

    def path(self, rel: str) -> Path:
        rel = rel.strip().replace("<network>", self.network)
        prefix = f"evidence/{self.network}/"
        if rel.startswith(prefix):
            return (self.net_dir / rel[len(prefix) :]).resolve()
        if rel == f"evidence/{self.network}":
            return self.net_dir
        if rel.startswith("evidence/"):
            rest = rel[len("evidence/") :]
            for r in self.roots:
                if (r / rest).exists():
                    return (r / rest).resolve()
            return (self.roots[0] / rest).resolve()
        return super().path(rel)


MISSING = "␠"  # stands in for a value whose evidence file is missing; never reaches a slide
# Values from optional evaluations: when their file is missing, the tiles, rows and sentences that use
# them are left out (also on Base Sepolia); every other value must resolve there.
OPTIONAL_PREFIXES = ("de_", "benign_")


def drop_missing(text: str) -> str:
    """Remove each sentence clause (split at "; ") or sentence that holds a missing value."""
    if MISSING not in text:
        return text
    out = []
    for sentence in re.split(r"(?<=[.!?])\s+", text):
        if MISSING not in sentence:
            out.append(sentence)
            continue
        end = sentence[-1] if sentence[-1] in ".!?" else ""
        clauses = [c for c in sentence.rstrip(".!?").split("; ") if MISSING not in c]
        if clauses:
            out.append("; ".join(clauses) + end)
    return " ".join(out).strip()


def plural(n: str | int | None, one: str, many: str) -> str:
    return one if str(n) == "1" else many


INLINE_RE = re.compile(r"\[`([^`]+)`\]\(([^)\s]+)\)|\[([^\]]+)\]\(([^)\s]+)\)|`([^`]+)`|\*\*([^*]+)\*\*")


class Values:
    """Resolves ``{{name}}`` / ``{{name|filter}}`` with the renderer; hash links become short form + URL."""

    def __init__(self, ctx: EvidenceContext, entries: dict[str, Any]) -> None:
        self.ctx, self.entries = ctx, entries
        self.missing: dict[str, str] = {}

    def raw(self, name: str) -> str | None:
        entry = self.entries.get(name)
        if entry is None:
            raise SystemExit(f"{{{{{name}}}}} is not in docs/deliverables/placeholders.md")
        try:
            return rd.resolve(self.ctx, entry)
        except rd.Unresolved as exc:
            self.missing[name] = str(exc)
            return None

    def get(self, name: str) -> str:
        v = self.raw(name)
        return "n/a" if v is None else v

    def url(self, value: str, prefer: str = "basescan") -> str | None:
        for name in (prefer, "basescan", "blockscout"):
            if name in self.ctx.explorers:
                return self.ctx.explorer_url(name, value)
        return None

    def short(self, value: str) -> str:
        return f"{value[:8]}...{value[-4:]}"

    def link_md(self, value: str, prefer: str = "basescan") -> str:
        url = self.url(value, prefer)
        return f"[`{self.short(value)}`]({url})" if url else f"`{self.short(value)}`"

    def render(self, text: str) -> str:
        """Render placeholders; a missing value becomes ``MISSING`` so the caller can drop that clause."""

        def sub(m: re.Match[str]) -> str:
            name, filt = m.group(1), m.group(2)
            value = self.raw(name)
            if value is None:
                return MISSING
            entry = self.entries[name]
            if filt in ("link", "basescan", "blockscout"):
                if entry.fmt not in ("hash", "address"):
                    raise SystemExit(f"|{filt} on {name}, which is not a hash or an address")
                return self.link_md(value, "blockscout" if filt == "blockscout" else "basescan")
            return rd.apply_filter(self.ctx, value, filt, entry)

        return rd.USE_RE.sub(sub, text)


def plain(md: str) -> str:
    """Markdown inline -> plain text (for notes and measuring)."""

    def sub(m: re.Match[str]) -> str:
        if m.group(1):
            return m.group(1)
        if m.group(3):
            return m.group(3)
        return m.group(5) or m.group(6) or ""

    return INLINE_RE.sub(sub, md)


def notes_text(md: str) -> str:
    def sub(m: re.Match[str]) -> str:
        if m.group(1):
            return f"{m.group(1)} ({m.group(2)})"
        if m.group(3):
            return f"{m.group(3)} ({m.group(4)})"
        return m.group(5) or m.group(6) or ""

    return INLINE_RE.sub(sub, md)


# ---------------------------------------------------------------------------------------------------
# deck.md parser
# ---------------------------------------------------------------------------------------------------

LABEL_RE = re.compile(r"^(On slide|Table|Visual|Notes \((\d\d:\d\d)\)|Sources):\s*(.*)$")


def parse_deck(text: str) -> dict[int, dict[str, Any]]:
    body = text.split("\n# Build notes", 1)[0]
    parts = re.split(r"^## Slide (\d+): (.*)$", body, flags=re.MULTILINE)
    slides: dict[int, dict[str, Any]] = {}
    for i in range(1, len(parts), 3):
        no, title, block = int(parts[i]), parts[i + 1].strip(), parts[i + 2]
        s: dict[str, Any] = {
            "title": title,
            "bullets": [],
            "table": [],
            "footnote": "",
            "notes_lines": [],
            "time": "",
            "sources": "",
            "visual": "",
        }
        label = None
        for line in block.splitlines():
            m = LABEL_RE.match(line)
            if m:
                label = m.group(1).split(" (")[0]
                if m.group(2):
                    s["time"] = m.group(2)
                rest = m.group(3).strip()
                if label == "Notes" and rest:
                    s["notes_lines"].append(rest)
                elif label == "Sources" and rest:
                    s["sources"] = rest
                elif label == "Visual" and rest:
                    s["visual"] = rest
                continue
            st = line.strip()
            if label == "On slide" and st.startswith("- "):
                s["bullets"].append(st[2:].strip())
            elif label == "Table" and st.startswith("|"):
                # a pipe inside {{name|filter}} is not a cell boundary
                cells = [c.strip() for c in re.split(r"\|(?![^{}]*\}\})", st.strip()[1:-1])]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    s["table"].append(cells)
            elif label == "Table" and st.startswith("Footnote under the table:"):
                s["footnote"] = st.split(":", 1)[1].strip()
            elif label == "Table" and st and s["footnote"] and not s["footnote"].endswith('"'):
                s["footnote"] += " " + st
            elif label == "Notes" and st:
                s["notes_lines"].append(st)
            elif label == "Sources" and st:
                s["sources"] += " " + st
            elif label == "Visual" and st:
                s["visual"] += " " + st
        slides[no] = s
    if sorted(slides) != list(range(1, 16)):
        raise SystemExit(f"deck.md: expected slides 1-15, found {sorted(slides)}")
    return slides


def notes_paragraphs(lines: list[str]) -> list[str]:
    paras: list[str] = []
    for line in lines:
        if not paras or line.startswith(("Say this line only if", "On compliance:")):
            paras.append(line)
        else:
            paras[-1] += " " + line
    return paras


# ---------------------------------------------------------------------------------------------------
# text fitting (Arial metrics from the system fonts; LibreOffice and PowerPoint both use Arial)
# ---------------------------------------------------------------------------------------------------


def _font_file(names: list[str]) -> str | None:
    for base in (Path("C:/Windows/Fonts"), Path("/usr/share/fonts/truetype/msttcorefonts"), Path("/Library/Fonts")):
        for n in names:
            if (base / n).is_file():
                return str(base / n)
    return None


_FONT_FILES = {
    (FONT, False): _font_file(["arial.ttf", "Arial.ttf"]),
    (FONT, True): _font_file(["arialbd.ttf", "Arial Bold.ttf"]),
    (MONO, False): _font_file(["cour.ttf", "Courier New.ttf"]),
    (MONO, True): _font_file(["courbd.ttf", "Courier New Bold.ttf"]),
}
_FONTS: dict[tuple[str, bool], Any] = {}


def text_width_pt(text: str, pt: float, bold: bool = False, font: str = FONT) -> float:
    key = (font, bold)
    if key not in _FONTS:
        f = _FONT_FILES.get(key)
        _FONTS[key] = ImageFont.truetype(f, 100) if f else None
    f = _FONTS[key]
    if f is None:
        return len(text) * pt * 0.55
    return f.getlength(text) * pt / 100


def wrapped_lines(text: str, pt: float, width_in: float, bold: bool = False, font: str = FONT) -> int:
    width_pt = width_in * 72
    lines = 0
    for raw in text.split("\n"):
        words, cur = raw.split(" "), ""
        n = 1
        for w in words:
            trial = (cur + " " + w).strip()
            if text_width_pt(trial, pt, bold, font) <= width_pt or not cur:
                cur = trial
            else:
                n, cur = n + 1, w
        lines += n
    return lines


WARNINGS: list[str] = []


def check_fit(label: str, paras: list[dict[str, Any]], width_in: float, height_in: float) -> None:
    need = 0.0
    for p in paras:
        size = p["size"]
        w = width_in - (p.get("indent", 0.0))
        n = wrapped_lines(plain(p["text"]), size, w, p.get("bold", False), p.get("font", FONT))
        need += n * size * 1.2 / 72 + p.get("space_after", 0) / 72
        for word in plain(p["text"]).split():
            if text_width_pt(word, size, p.get("bold", False), p.get("font", FONT)) > w * 72:
                WARNINGS.append(f"{label}: word {word!r} is wider than its box at {size} pt")
    if need > height_in + 0.1:  # one line's leading of slack: text boxes do not clip
        WARNINGS.append(f"{label}: text needs {need:.2f} in, box has {height_in:.2f} in")


# ---------------------------------------------------------------------------------------------------
# drawing helpers
# ---------------------------------------------------------------------------------------------------


def _para_spec(p: Any, defaults: dict[str, Any]) -> dict[str, Any]:
    spec = dict(defaults)
    spec.update(p if isinstance(p, dict) else {"text": p})
    return spec


def set_bullet(para, indent: float = 0.28, color: str = ACCENT) -> None:
    pPr = para._p.get_or_add_pPr()
    pPr.set("marL", str(int(In(indent))))
    pPr.set("indent", str(-int(In(indent))))
    clr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(clr, qn("a:srgbClr")).set("val", color)
    etree.SubElement(pPr, qn("a:buFont")).set("typeface", FONT)
    etree.SubElement(pPr, qn("a:buChar")).set("char", "\u2022")


def add_runs(para, md: str, size: float, bold: bool, color: str, font: str) -> None:
    pos = 0

    def run(text: str, *, b: bool = bold, f: str = font, c: str = color, link: str | None = None) -> None:
        if not text:
            return
        r = para.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = b
        r.font.name = f
        r.font.color.rgb = rgb(c)
        if link:
            r.hyperlink.address = link
            r.font.underline = True

    link_color = OI_BLUE if color in (INK, MUTED, ACCENT_DARK) else OI_SKY if color == WHITE else color
    for m in INLINE_RE.finditer(md):
        run(md[pos : m.start()])
        if m.group(1):
            run(m.group(1), link=m.group(2), c=link_color)
        elif m.group(3):
            run(m.group(3), link=m.group(4), c=link_color)
        elif m.group(5):
            run(m.group(5), f=MONO)
        elif m.group(6):
            run(m.group(6), b=True)
        pos = m.end()
    run(md[pos:])


def fill_frame(
    tf,
    paras: list[Any],
    *,
    size: float = 18,
    color: str = INK,
    bold: bool = False,
    align: str = "l",
    anchor: str = "t",
    font: str = FONT,
    bullet: bool = False,
    space_after: float = 6,
    margin: float = 0.08,
    wrap: bool = True,
) -> list[dict[str, Any]]:
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = In(margin)
    tf.margin_top = tf.margin_bottom = In(0.04)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    defaults = {
        "size": size,
        "color": color,
        "bold": bold,
        "align": align,
        "font": font,
        "bullet": bullet,
        "space_after": space_after,
    }
    specs = [_para_spec(p, defaults) for p in paras]
    for i, spec in enumerate(specs):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        for r in list(para.runs):
            r._r.getparent().remove(r._r)
        para.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[spec["align"]]
        para.space_after = Pt(spec["space_after"]) if i < len(specs) - 1 else Pt(0)
        para.line_spacing = 1.0
        if spec["bullet"]:
            set_bullet(para)
            spec["indent"] = 0.28
        add_runs(para, spec["text"], spec["size"], spec["bold"], spec["color"], spec["font"])
    return specs


def textbox(slide, x, y, w, h, paras, *, label: str = "", check: bool = True, name: str | None = None, **kw):
    tb = slide.shapes.add_textbox(In(x), In(y), In(w), In(h))
    if name:
        tb.name = name
    specs = fill_frame(tb.text_frame, paras if isinstance(paras, list) else [paras], **kw)
    margin = kw.get("margin", 0.08)
    if check:
        check_fit(label or plain(specs[0]["text"])[:40], specs, w - 2 * margin, h - 0.08)
    return tb


def box(
    slide,
    x,
    y,
    w,
    h,
    *,
    fill: str | None = CARD,
    line: str | None = CARD_LINE,
    line_w: float = 1.0,
    radius: float = 0.06,
    dash: bool = False,
    shape=MSO_SHAPE.ROUNDED_RECTANGLE,
    alpha: int | None = None,
    paras=None,
    label: str = "",
    check: bool = True,
    name: str | None = None,
    **kw,
):
    shp = slide.shapes.add_shape(shape, In(x), In(y), In(w), In(h))
    if name:
        shp.name = name
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = radius
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
        if alpha is not None:
            srgb = shp.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
            etree.SubElement(srgb, qn("a:alpha")).set("val", str(alpha * 1000))
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(line_w)
        if dash:
            shp.line.dash_style = MSO_LINE.DASH
    shp.shadow.inherit = False
    if paras is not None:
        margin = kw.get("margin", 0.08)
        specs = fill_frame(shp.text_frame, paras if isinstance(paras, list) else [paras], **kw)
        if check:
            check_fit(label or plain(specs[0]["text"])[:40], specs, w - 2 * margin, h - 0.08)
    return shp


def arrow(slide, pts: list[tuple[float, float]], color: str = INK, width: float = 1.75, dash: bool = False):
    if len(pts) == 2:
        (x1, y1), (x2, y2) = pts
        shp = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, In(x1), In(y1), In(x2), In(y2))
    else:
        fb = slide.shapes.build_freeform(int(In(pts[0][0])), int(In(pts[0][1])), scale=1.0)
        fb.add_line_segments([(int(In(x)), int(In(y))) for x, y in pts[1:]], close=False)
        shp = fb.convert_to_shape()
        shp.fill.background()
        shp.shadow.inherit = False
    shp.line.color.rgb = rgb(color)
    shp.line.width = Pt(width)
    if dash:
        shp.line.dash_style = MSO_LINE.DASH
    ln = shp.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    return shp


def table(
    slide,
    x,
    y,
    widths,
    rows: list[list[str]],
    *,
    size: float = 16,
    header_size: float | None = None,
    row_h: float = 0.4,
    header_h: float | None = None,
    header_fill: str = ACCENT_DARK,
    fills: dict[int, str] | None = None,
    bold_rows: set[int] | None = None,
    aligns: list[str] | None = None,
    merges: list[tuple[int, int, int]] | None = None,
    label: str = "table",
    font_cols: dict[int, str] | None = None,
):
    nrows, ncols = len(rows), len(widths)
    gf = slide.shapes.add_table(nrows, ncols, In(x), In(y), In(sum(widths)), In(row_h * nrows))
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    for attr in ("bandRow", "firstRow"):
        tblPr.set(attr, "0")
    for c, wd in enumerate(widths):
        tbl.columns[c].width = In(wd)
    for r in range(nrows):
        tbl.rows[r].height = In(header_h if (r == 0 and header_h) else row_h)
    for r, cells in enumerate(rows):
        for c in range(ncols):
            cell = tbl.cell(r, c)
            text = cells[c] if c < len(cells) else ""
            cell.margin_left = cell.margin_right = In(0.06)
            cell.margin_top = cell.margin_bottom = In(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = rgb(header_fill)
            else:
                cell.fill.fore_color.rgb = rgb((fills or {}).get(r, WHITE if r % 2 else "F7F8FC"))
            tf = cell.text_frame
            tf.word_wrap = True
            align = (aligns or ["l"] * ncols)[c]
            sz = (header_size or size) if r == 0 else size
            bold = r == 0 or r in (bold_rows or set())
            for k, line in enumerate(text.split("\n")):  # "\n" starts a new line in the cell (one hash per line)
                para = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                para.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
                add_runs(para, line, sz, bold, WHITE if r == 0 else INK, (font_cols or {}).get(c, FONT))
            for word in plain(text).split():
                if text_width_pt(word, sz, bold) > (widths[c] - 0.12) * 72:
                    WARNINGS.append(f"{label}: word {word!r} (row {r}, col {c}) is wider than its column")
    for r, c1, c2 in merges or []:
        tbl.cell(r, c1).merge(tbl.cell(r, c2))
        tbl.cell(r, c1).text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
    return gf


def table_height(
    widths: list[float], rows: list[list[str]], size: float, row_h: float, header_h: float | None = None
) -> float:
    """Estimated rendered height: a row grows to its tallest wrapped cell (cell margins as in ``table``)."""
    total = 0.0
    for r, cells in enumerate(rows):
        lines = max(wrapped_lines(plain(c), size, w - 0.12, r == 0) for c, w in zip(cells, widths))
        total += max(header_h if (r == 0 and header_h) else row_h, lines * size * 1.2 / 72 + 0.1)
    return total


def title(slide, text: str) -> None:
    size = 30 if text_width_pt(text, 30, True) <= 9.9 * 72 else 24
    textbox(
        slide,
        0.42,
        0.1,
        10.2,
        0.98,
        [text],
        size=size,
        bold=True,
        color=WHITE,
        anchor="m",
        label=f"title {text[:30]}",
        name="Title",
    )


def footer(slide, md: str, size: float = 11) -> None:
    textbox(
        slide,
        0.5,
        6.86,
        12.35,
        0.38,
        [md],
        size=size,
        color=MUTED,
        anchor="t",
        label="footer",
        name="Footer",
        space_after=0,
    )


def notes(slide, md: str) -> None:
    slide.notes_slide.notes_text_frame.text = notes_text(md)


def picture(slide, path: Path, x, y, w=None, h=None):
    return slide.shapes.add_picture(str(path), In(x), In(y), In(w) if w else None, In(h) if h else None)


# ---------------------------------------------------------------------------------------------------
# generated assets
# ---------------------------------------------------------------------------------------------------


def tier_chart(path: Path, timeline: str, hashes: dict[str, str]) -> None:
    """Step chart of the wallet tier from the rendered ``wallet_tier_timeline`` value."""
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    level = {"FROZEN": -1, "0": 0, "1": 1, "2": 2, "3": 3}
    parts = [p.strip() for p in timeline.split("->")]
    start = parts[0]
    changes = []
    for p in parts[1:]:
        m = re.match(r"^(\S+)\s*\((?:before )?(S\d+)(?:,\s*([\d, ]+))?\)$", p)
        if not m:
            raise SystemExit(f"tier timeline part not understood: {p!r}")
        changes.append((m.group(1), int(m.group(2)[1:]), (m.group(3) or "").strip()))
    x0, x1 = 6, max([18] + [c[1] for c in changes])
    xs, ys = [x0], [level[start]]
    for lab, step, _ in changes:
        xs += [step, step]
        ys += [ys[-1], level[lab]]
    xs.append(x1 + 0.6)
    ys.append(ys[-1])
    # drawn at the size it is placed on the slide (6.2 in wide), so 14 pt here is 14 pt on the slide;
    # the transaction hashes are listed under the chart, not in it
    fig, ax = plt.subplots(figsize=(6.2, 2.45), dpi=220)
    ax.set_title("Wallet tier through the run (step, reason code)", fontsize=14, loc="left", color="#1E2235")
    ax.plot(xs, ys, color="#0072B2", linewidth=3, solid_joinstyle="miter")
    prev = level[start]
    for lab, step, codes in changes:
        y = level[lab]
        color = "#D55E00" if lab == "FROZEN" else "#0072B2"
        ax.plot([step], [y], marker="s" if lab == "FROZEN" else "o", color=color, markersize=8, zorder=3)
        text = f"S{step}" + (f" ({codes})" if codes else "")
        # a rise is labelled above right, a fall below left (clear of the next vertical segment), a freeze
        # below right
        up = y >= prev
        right = up or lab == "FROZEN"
        ax.annotate(
            text,
            (step, y),
            xytext=(step + (0.25 if right else -0.25), y + (0.3 if up else -0.3)),
            fontsize=14,
            va="bottom" if up else "top",
            ha="left" if right else "right",
            color="#1E2235",
        )
        prev = y
    ax.set_yticks([-1, 0, 1, 2, 3])
    ax.set_yticklabels(["FROZEN", "0", "1", "2", "3"], fontsize=14)
    ticks = list(range(x0, x1 + 1, 2))
    ax.set_xticks(ticks)
    ax.set_xticklabels([f"S{t}" for t in ticks], fontsize=14)
    ax.set_ylim(-2.3, 3.7)
    ax.set_xlim(x0 - 0.3, x1 + 3.4)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.grid(axis="y", color="#E3E6F0", linewidth=0.8)
    fig.tight_layout(pad=0.3)
    fig.savefig(path, transparent=False, facecolor="white")
    plt.close(fig)


def qr_png(path: Path, url: str) -> bool:
    try:
        import qrcode
    except ImportError:
        WARNINGS.append("qrcode not installed: thank-you slide built without a QR code (add --with qrcode)")
        return False
    img = qrcode.make(url, box_size=12, border=2)
    img.save(str(path))
    return True


# ---------------------------------------------------------------------------------------------------
# the slides
# ---------------------------------------------------------------------------------------------------


class Builder:
    def __init__(
        self, prs, slides: dict[int, dict[str, Any]], vals: Values, network: str, draft: bool, blockscout_verified: bool
    ) -> None:
        self.prs, self.S, self.v = prs, slides, vals
        self.network, self.draft = network, draft
        self.blockscout_verified = blockscout_verified
        self.net_label = "Base Sepolia" if network == "base_sepolia" else f"a local {network.capitalize()} dry run"
        self.layout = prs.slide_layouts[1]  # 标题和内容: white page with the FinTechathon banner
        self.content: list[Any] = []

    # common parts --------------------------------------------------------------------------------

    def r(self, md: str) -> str:
        return self.v.render(md)

    def bullets(self, n: int) -> list[str]:
        """The slide's bullets, rendered, by position; a bullet whose only content was missing is ""."""
        return [drop_missing(self.r(b)) for b in self.S[n]["bullets"]]

    def have(self, *names: str) -> bool:
        return all(self.v.raw(n) is not None for n in names)

    def speaker_notes(self, slide, n: int, paras: list[str] | None = None, extra: list[str] | None = None) -> None:
        """Spoken text only: the deck.md "Sources" lines stay in deck.md, out of the presenter view."""
        s = self.S[n]
        body = paras if paras is not None else notes_paragraphs(s["notes_lines"])
        spoken = [t for t in (drop_missing(self.r(p)) for p in body + (extra or [])) if t]
        text = f"[{s['time']}] " + "\n".join(spoken)
        if self.draft:
            text += f"\n\nDRAFT: values from {self.net_label}; rebuild on the Base Sepolia evidence before presenting."
        notes(slide, text)

    def draft_tag(self, slide, dark: bool = False) -> None:
        if not self.draft:
            return
        chain = self.v.get("chain_id")
        text = f"DRAFT: numbers from {self.net_label} (chain id {chain}), not the recorded Base Sepolia run"
        textbox(
            slide,
            5.2,
            7.18,
            7.95,
            0.3,
            [text],
            size=10,
            color=OI_YELLOW if dark else DRAFT_ON_WHITE,
            align="r",
            label="draft tag",
            name="Draft tag",
            bold=True,
        )

    def new_slide(self, n: int, reuse=None):
        slide = reuse or self.prs.slides.add_slide(self.layout)
        title(slide, self.S[n]["title"])
        self.draft_tag(slide)
        return slide

    # 1 ------------------------------------------------------------------------------------------

    def slide1(self) -> None:
        slide = self.prs.slides[0]
        b = self.bullets(1)
        shapes = {sh.name: sh for sh in slide.shapes}
        name_box, team_box = shapes["文本框 1"], shapes["文本框 2"]
        panel = box(slide, 0.55, 4.4, 8.45, 2.75, fill="0B0E26", line=None, radius=0.04, alpha=55)
        slide.shapes._spTree.remove(panel._element)
        slide.shapes._spTree.insert(2, panel._element)  # behind the template's text boxes
        for sh, y, h in ((name_box, 4.45, 0.7), (team_box, 5.1, 0.45)):
            sh.left, sh.top, sh.width, sh.height = In(0.74), In(y), In(8.0), In(h)
        fill_frame(name_box.text_frame, [b[0]], size=40, bold=True, color=WHITE, wrap=False)
        fill_frame(team_box.text_frame, [b[1]], size=20, bold=True, color=WHITE, wrap=False)
        members = b[3].split(" and ", 1)
        event, topic = (b[4].split(", Topic F", 1) + [""])[:2]
        textbox(
            slide,
            0.74,
            5.55,
            8.1,
            1.3,
            [
                {"text": b[2], "bold": True, "size": 16},
                members[0] + " and",
                members[1],
                event + ("," if topic else ""),
                "Topic F" + topic if topic else "",
            ],
            size=14,
            color=WHITE,
            space_after=1,
            label="title sub",
        )
        chain = self.v.get("chain_id")
        if self.network == "base_sepolia":
            foot = f"Recorded run on Base Sepolia (chain id {chain}). Code and evidence: github.com/taivnaaiwnl/shenzhen-comp"
        else:
            foot = f"Local dry run (chain id {chain}). Code and evidence: github.com/taivnaaiwnl/shenzhen-comp"
        textbox(slide, 0.74, 6.87, 8.6, 0.3, [foot], size=12, color="C9CEE6", label="title footer")
        self.draft_tag(slide, dark=True)
        self.speaker_notes(slide, 1)

    # 2 ------------------------------------------------------------------------------------------

    def slide2(self) -> None:
        slide = self.prs.slides[1]
        shapes = {sh.name: sh for sh in slide.shapes}
        shapes["文本框 7"]._element.getparent().remove(shapes["文本框 7"]._element)  # Chinese "contents"
        head = shapes["文本框 4"]
        head.left, head.top, head.width, head.height = In(0.9), In(3.3), In(3.0), In(0.8)
        fill_frame(head.text_frame, ["CONTENTS"], size=36, bold=True, color=WHITE, wrap=False)
        items = [re.match(r"^(\d\d)\s+(.*)$", b).groups() for b in self.bullets(2)]
        pairs = [
            ("文本框 5", "文本框 6"),
            ("文本框 8", "文本框 9"),
            ("文本框 10", "文本框 11"),
            ("文本框 12", "文本框 13"),
            ("文本框 14", "文本框 15"),
            ("文本框 16", "文本框 17"),
        ]
        cols, rows = (4.3, 8.55), (2.3, 3.45, 4.6)
        for i, ((num_name, txt_name), (num, text)) in enumerate(zip(pairs, items)):
            x, y = cols[i % 2], rows[i // 2]
            nb, tb = shapes[num_name], shapes[txt_name]
            nb.left, nb.top, nb.width, nb.height = In(x), In(y), In(0.85), In(0.62)
            tb.left, tb.top, tb.width, tb.height = In(x + 0.85), In(y), In(3.3), In(0.95)
            fill_frame(nb.text_frame, [num], size=30, bold=True, color=ACCENT, wrap=False)
            fill_frame(tb.text_frame, [text], size=24, color=WHITE)
            check_fit(f"contents {num}", [{"text": text, "size": 24}], 3.3 - 0.16, 0.95)
        self.draft_tag(slide, dark=True)
        self.speaker_notes(slide, 2)

    # 3 ------------------------------------------------------------------------------------------

    def slide3(self) -> None:
        slide = self.prs.slides[2]
        for sh in list(slide.shapes):  # the template's Chinese main-title box; the title is redrawn below
            sh._element.getparent().remove(sh._element)
        self.new_slide(3, reuse=slide)
        b = self.bullets(3)
        textbox(slide, 0.6, 1.38, 12.1, 0.5, [b[0]], size=20, color=INK, bold=True, label="s3 lead")
        causes = ["persuaded", "injected", "no caps"]
        for i, (text, cause) in enumerate(zip(b[1:4], causes)):
            y = 2.0 + i * 1.57
            name, rest = text.split(" (", 1)
            meta, detail = rest.split("): ", 1)
            box(
                slide,
                0.6,
                y,
                7.3,
                1.42,
                paras=[
                    {"text": name, "bold": True, "size": 20},
                    {"text": meta, "size": 16, "color": MUTED},
                    {"text": detail[0].upper() + detail[1:], "size": 18},
                ],
                label=f"s3 card {i}",
                space_after=2,
                margin=0.15,
                anchor="m",
            )
            box(
                slide,
                6.1,
                y + 0.15,
                1.6,
                0.42,
                fill=OI_VERM,
                line=None,
                radius=0.5,
                paras=[{"text": cause, "size": 15, "bold": True, "color": WHITE, "align": "c"}],
                anchor="m",
                label="cause chip",
                margin=0.02,
            )
        box(slide, 8.2, 2.0, 4.53, 4.56, fill=CARD, line=CARD_LINE, radius=0.04)
        textbox(
            slide,
            8.4,
            2.12,
            4.2,
            0.5,
            [{"text": "Know Your Agent (KYA)", "bold": True, "size": 20, "color": ACCENT_DARK}],
            label="kya head",
        )
        qs = [
            ("Who is the agent?", True),
            ("Who authorised it?", True),
            ("Which permissions?", True),
            ("How has it behaved when paid?", False),
        ]
        for i, (q, done) in enumerate(qs):
            y = 2.75 + i * 0.72
            if done:
                box(
                    slide,
                    8.45,
                    y + 0.1,
                    0.36,
                    0.36,
                    fill=OI_BLUE,
                    line=None,
                    radius=0.15,
                    paras=[{"text": "\u2713", "size": 16, "bold": True, "color": WHITE, "align": "c"}],
                    anchor="m",
                    margin=0,
                    check=False,
                )
            else:
                box(slide, 8.45, y + 0.1, 0.36, 0.36, fill=WHITE, line=OI_VERM, line_w=2.25, radius=0.15, dash=True)
            textbox(
                slide,
                8.95,
                y - 0.05,
                3.7,
                0.66,
                [{"text": q, "size": 18, "bold": not done, "color": INK if done else OI_VERM}],
                anchor="m",
                label=f"kya {i}",
            )
        textbox(slide, 8.4, 5.68, 4.2, 0.8, [{"text": b[4], "size": 16, "color": MUTED}], label="kya foot")
        self.speaker_notes(slide, 3)

    # 4 ------------------------------------------------------------------------------------------

    def slide4(self) -> None:
        slide = self.new_slide(4)
        b = self.bullets(4)
        heads = ["Credit bureau", "Grade gate", "Earned Autonomy", "LLM is not the boundary"]
        for i, (h, text) in enumerate(zip(heads, b)):
            x, y = 0.6 + (i % 2) * 6.15, 1.4 + (i // 2) * 1.38
            box(
                slide,
                x,
                y,
                5.95,
                1.25,
                paras=[{"text": h, "bold": True, "size": 20, "color": ACCENT_DARK}, {"text": text, "size": 18}],
                space_after=2,
                margin=0.15,
                label=f"s4 card {i}",
            )
        textbox(
            slide,
            0.6,
            4.2,
            2.1,
            0.45,
            [{"text": "Topic F:", "bold": True, "size": 16, "color": MUTED}],
            anchor="m",
            label="topic f",
        )
        for i, word in enumerate(["asset management", "transaction execution", "risk control"]):
            box(
                slide,
                2.0 + i * 3.0,
                4.22,
                2.8,
                0.42,
                fill=WHITE,
                line=ACCENT,
                radius=0.5,
                paras=[{"text": word, "size": 16, "color": ACCENT_DARK, "align": "c"}],
                anchor="m",
                margin=0.02,
                label="topic chip",
            )
        failed = self.v.raw("atk_failed")
        tiles = [
            (self.r("{{tx_total}}"), self.r("on-chain transactions ({{tx_agent_key}} agent-signed)")),
            (
                self.r("{{atk_failed}} of {{atk_total}}"),
                "attacks that reached their goal" if failed != "1" else "attack that reached its goal",
            ),
            (self.r("{{atk_funds_lost_usdc}}"), "USDC lost to attacks"),
            (self.r("{{grade_changes_recompute_match}} / {{grade_changes_total}}"), "grades recompute to MATCH"),
        ]
        if self.have("benign_false_positives", "benign_tasks"):
            tiles.insert(
                3, (self.r("{{benign_false_positives}} / {{benign_tasks}}"), "false positives on normal tasks")
            )
        n = len(tiles)
        w = (12.13 - (n - 1) * 0.2) / n
        for i, (big, small) in enumerate(tiles):
            x = 0.6 + i * (w + 0.2)
            box(
                slide,
                x,
                4.9,
                w,
                1.85,
                fill=CARD,
                line=CARD_LINE,
                paras=[
                    {"text": big, "size": 30, "bold": True, "color": OI_BLUE, "align": "c"},
                    {"text": small, "size": 16, "color": INK, "align": "c"},
                ],
                space_after=4,
                margin=0.1,
                anchor="m",
                label=f"s4 tile {i}",
            )
        # deck.md says "{{atk_failed}} failed attacks"; say it so it reads one way and agrees in number
        phrase = (
            f"{failed} {plural(failed, 'attack that reached its goal', 'attacks that reached their goal')}"
            if failed not in (None, "0")
            else "no attack that reached its goal"
        )
        paras = [p.replace("{{atk_failed}} failed attacks", phrase) for p in notes_paragraphs(self.S[4]["notes_lines"])]
        self.speaker_notes(slide, 4, paras)

    # 5 ------------------------------------------------------------------------------------------

    def slide5(self) -> None:
        slide = self.new_slide(5)
        b = self.bullets(5)
        # zones
        box(slide, 0.5, 1.38, 8.05, 2.55, fill="F7F8FC", line=CARD_LINE, radius=0.03)
        textbox(
            slide,
            0.6,
            1.4,
            4.0,
            0.34,
            [{"text": "Off-chain processes, one key each", "size": 13, "bold": True, "color": MUTED}],
            label="zone off",
        )
        box(slide, 0.5, 4.08, 8.05, 2.7, fill="EEF6FB", line=CARD_LINE, radius=0.03)
        # bottom right of the zone, clear of the Bureau -> ERC-8004 connector on the left
        textbox(
            slide,
            5.95,
            6.3,
            2.5,
            0.4,
            [{"text": "On-chain (Base Sepolia)", "size": 14, "bold": True, "color": MUTED, "align": "r"}],
            label="zone on",
        )

        def comp(x, y, w, h, name, sub, fill=WHITE, line=ACCENT, size=15):
            box(
                slide,
                x,
                y,
                w,
                h,
                fill=fill,
                line=line,
                line_w=1.5,
                radius=0.1,
                paras=[
                    {"text": name, "bold": True, "size": size, "align": "c"},
                    {"text": sub, "size": 13, "color": MUTED, "align": "c"},
                ],
                anchor="m",
                space_after=0,
                margin=0.04,
                label=f"comp {name}",
            )

        # wider gaps between the columns, so the arrow labels sit clear of the dashed "assume hijacked" box
        C1, C2, C3, W1, W2, W3 = 0.8, 3.55, 6.6, 2.05, 2.2, 1.85
        comp(C1, 1.82, W1, 0.8, "Owner console", "owner key")
        comp(C2, 1.82, W2, 0.8, "Wallet agent + LLM", "no keys")
        comp(C3, 1.82, W3, 0.8, "Counterparties", "simulated, own keys")
        comp(C1, 2.95, W1, 0.8, "Bureau", "bureau key")
        comp(C2, 2.95, W2, 0.8, "Signer", "the only agent key")
        comp(C3, 2.95, W3, 0.8, "Autonomy operator", "autonomy key")
        comp(C1, 4.55, W1, 0.85, "AgentCreditBureau", "grades, evidence hashes", fill="FFFFFF", line=OI_BLUE)
        comp(C2, 4.55, W2, 0.85, "TierVault (wallet)", "mandate, caps, grade gate", fill="FFFFFF", line=OI_BLUE)
        comp(C3, 4.55, W3, 0.85, "USDC", "test token", fill="FFFFFF", line=OI_BLUE)
        comp(
            C1,
            5.72,
            C2 + W2 - C1,
            0.85,
            "ERC-8004 Identity and Reputation",
            "official registries, Draft standard",
            fill="FFFFFF",
            line=OI_BLUE,
        )
        # assume hijacked
        box(slide, C2 - 0.1, 1.72, W2 + 0.2, 1.0, fill=None, line=OI_VERM, line_w=2.25, dash=True, radius=0.12)
        textbox(
            slide,
            C2 - 0.1,
            1.42,
            W2 + 0.2,
            0.3,
            [{"text": "assume hijacked", "size": 12, "bold": True, "color": OI_VERM, "align": "r"}],
            label="hijacked",
        )
        k = INK
        arrow(slide, [(C1 + W1, 2.22), (C2 - 0.1, 2.22)], k)
        arrow(slide, [(C3, 2.22), (C2 + W2 + 0.1, 2.22)], k)
        arrow(slide, [(C2 + W2 / 2, 2.72), (C2 + W2 / 2, 2.95)], k)
        arrow(slide, [(C2 + W2 / 2, 3.75), (C2 + W2 / 2, 4.55)], k)
        arrow(slide, [(C2, 4.97), (C1 + W1, 4.97)], k)
        arrow(slide, [(C2 + W2, 4.97), (C3, 4.97)], k)
        arrow(slide, [(C2 + W2 / 2 + 0.6, 5.4), (C2 + W2 / 2 + 0.6, 5.72)], k)
        arrow(slide, [(C1 + W1 / 2 + 0.3, 3.75), (C1 + W1 / 2 + 0.3, 4.55)], k)
        arrow(slide, [(C1 + 0.2, 3.75), (C1 + 0.2, 3.95), (0.64, 3.95), (0.64, 6.14), (C1, 6.14)], k)
        arrow(slide, [(C3 + 0.5, 3.75), (C3 + 0.5, 4.3), (C2 + W2 - 0.4, 4.3), (C2 + W2 - 0.4, 4.55)], k)
        g1, g2 = (C2 - 0.1) - (C1 + W1), C3 - (C2 + W2 + 0.1)  # free space beside the dashed box
        labels = [
            (C1 + W1, 1.86, g1, "intent", "c"),
            (C2 + W2 + 0.1, 1.86, g2, "invoice", "c"),
            (C2 + W2 / 2 + 0.05, 3.76, 1.3, "signed tx", "l"),
            (C1 + W1, 4.62, C2 - C1 - W1, "grade", "c"),
            (C2 + W2, 4.62, C3 - C2 - W2, "wallet", "c"),
            (C1 + W1 / 2 + 0.35, 3.76, 0.8, "posts", "l"),
            (C3 + 0.55, 3.76, 1.0, "tier", "l"),
            (C2 + W2 / 2 + 0.65, 5.4, 1.0, "reads", "l"),
        ]
        for x, y, w, t, al in labels:
            textbox(
                slide,
                x,
                y,
                w,
                0.32,
                [{"text": t, "size": 13, "color": MUTED, "align": al}],
                margin=0.0,
                label=f"arrow {t}",
            )
        textbox(slide, 8.75, 1.38, 4.0, 5.45, b, size=18, bullet=True, space_after=7, label="s5 bullets")
        acb, vault = self.r("{{addr_acb|short}}"), self.r("{{addr_vault_wallet|short}}")
        ident, rep = self.r("{{addr_identity_registry|short}}"), self.r("{{addr_reputation_registry|short}}")
        footer(
            slide,
            f"Addresses: ACB `{acb}`, wallet vault `{vault}`, Identity `{ident}`, Reputation `{rep}`"
            + ("" if self.network == "base_sepolia" else f" ({self.network})"),
        )
        self.speaker_notes(slide, 5)

    # 6 ------------------------------------------------------------------------------------------

    def slide6(self) -> None:
        slide = self.new_slide(6)
        b = self.bullets(6)
        textbox(slide, 0.5, 1.38, 6.1, 3.95, b, size=18, bullet=True, space_after=8, label="s6 bullets")
        tx = self.v.raw("tx_s11_grade_ghost") or "0x" + "0" * 64
        box(
            slide,
            0.6,
            5.42,
            5.9,
            0.85,
            fill="1E2235",
            line=None,
            radius=0.08,
            paras=[
                {
                    "text": f"$ acb bureau recompute --tx {self.v.short(tx)}",
                    "font": MONO,
                    "size": 14,
                    "color": "E6E8F2",
                },
                {"text": self.r("{{recompute_ghost_s11}}"), "font": MONO, "size": 14, "bold": True, "color": OI_GREEN},
            ],
            space_after=2,
            margin=0.15,
            anchor="m",
            label="cmd",
        )
        # Ghost Feeds evidence card
        sc = yaml.safe_load((REPO / "config" / "scorecard.yaml").read_text(encoding="utf-8"))
        base, (lo, hi), bands = sc["base"], sc["clamp"], sc["bands"]
        ev = self.v.ctx.evidence()
        sel = "gradeChanges[step=S11,agent=ghost][-1]"
        bundle_rel = rd.evaluate(ev, sel + ".bundle")
        bundle = self.v.ctx.load(f"evidence/<network>/{bundle_rel}")
        points = bundle["points"]
        score = int(self.v.get("score_ghost_s11"))
        if max(lo, min(hi, base + sum(points.values()))) != score:
            WARNINGS.append("slide 6: bundle points do not add up to the posted score")
        X, Y, W = 6.85, 1.38, 5.9
        box(slide, X, Y, W, 5.4, fill=CARD, line=CARD_LINE, radius=0.03)
        textbox(
            slide,
            X + 0.15,
            Y + 0.05,
            W - 0.3,
            0.45,
            [{"text": "Ghost Feeds: grade posting at S11", "bold": True, "size": 18, "color": ACCENT_DARK}],
            label="ghost head",
        )
        sx, sw = X + 2.15, W - 2.45  # scale area

        def px(v: float) -> float:
            return sx + (v - lo) / (hi - lo) * sw

        by = Y + 0.6
        edges = [
            (lo, bands["C"], "D", OI_VERM),
            (bands["C"], bands["B"], "C", OI_ORANGE),
            (bands["B"], bands["A"], "B", OI_SKY),
            (bands["A"], hi, "A", OI_GREEN),
        ]
        textbox(
            slide,
            X + 0.15,
            by,
            2.0,
            0.38,
            [{"text": "Grade bands", "size": 14, "color": MUTED}],
            anchor="m",
            label="bands label",
        )
        for a, z, letter, col in edges:
            box(
                slide,
                px(a),
                by,
                px(z) - px(a),
                0.38,
                fill=col,
                line=WHITE,
                shape=MSO_SHAPE.RECTANGLE,
                paras=[
                    {"text": letter, "size": 14, "bold": True, "color": WHITE if letter != "C" else INK, "align": "c"}
                ],
                anchor="m",
                margin=0,
                label=f"band {letter}",
            )
        for t in (lo, bands["C"], bands["B"], bands["A"], hi):
            textbox(
                slide,
                px(t) - 0.3,
                by + 0.38,
                0.6,
                0.26,
                [{"text": str(t), "size": 10, "color": MUTED, "align": "c"}],
                margin=0,
                label="tick",
            )
        names = {
            "F1": "F1 fulfilment",
            "F2": "F2 failures",
            "F3": "F3 flags",
            "F4a": "F4a experience",
            "F4b": "F4b age",
            "F5a": "F5a diversity",
            "F5b": "F5b partner quality",
            "F5c": "F5c circular volume",
            "F6": "F6 volume spike",
        }
        rows = [("Base", None, base)]
        cur = base
        for f, p in points.items():
            if p:
                rows.append((f"{names.get(f, f)} {p:+d}", cur, cur + p))
                cur += p
        zero = [f for f, p in points.items() if not p]
        rows.append((f"Score {score}", None, score))
        ry = by + 0.8
        for i, (lab, a, z) in enumerate(rows):
            y = ry + i * 0.42
            textbox(
                slide,
                X + 0.15,
                y,
                2.0,
                0.36,
                [{"text": lab, "size": 14, "bold": a is None}],
                anchor="m",
                margin=0.02,
                label=f"wf {lab}",
            )
            if a is None:
                box(slide, px(z) - 0.07, y + 0.05, 0.14, 0.26, fill=INK, line=None, shape=MSO_SHAPE.RECTANGLE)
            else:
                left, right = sorted((px(a), px(z)))
                col = OI_VERM if z < a else OI_BLUE
                box(
                    slide, left, y + 0.06, max(right - left, 0.03), 0.24, fill=col, line=None, shape=MSO_SHAPE.RECTANGLE
                )
        y = ry + len(rows) * 0.42 + 0.02
        textbox(
            slide,
            X + 0.15,
            y,
            W - 0.3,
            0.3,
            [{"text": f"0 points: {', '.join(zero)}", "size": 12, "color": MUTED}],
            label="zero points",
        )
        reasons = [int(c) for c in rd.evaluate(ev, sel + ".reasons")]
        rl = [{"text": "Reason codes", "size": 14, "bold": True}] + [
            {"text": f"{c}  {reason_codes.text(c)}", "size": 14} for c in reasons[:4]
        ]
        textbox(slide, X + 0.15, y + 0.32, W - 0.3, 1.45, rl, space_after=1, label="reasons")
        footer(
            slide,
            self.r(
                "Ghost Feeds: {{grade_path_ghost}}; score {{score_ghost_s11}}; reasons "
                "{{reasons_ghost_s11}}; recompute {{recompute_ghost_s11}}; grade tx "
                "{{tx_s11_grade_ghost|link}}. Bands and base: config/scorecard.yaml"
            ),
            size=12,
        )
        self.speaker_notes(slide, 6)

    # 7 ------------------------------------------------------------------------------------------

    def slide7(self) -> None:
        slide = self.new_slide(7)
        s = self.S[7]
        rows = [[self.r(c) for c in r] for r in s["table"]]
        rows[0] = (
            ["Tier", "Name", "Per payment", "24 h window", "Tx per window", "Pays directly", "Clean orders to promote"]
            if len(rows[0]) == 7
            else rows[0]
        )
        frozen = next((i for i, r in enumerate(rows) if r[0] == "FROZEN"), None)
        table(
            slide,
            0.6,
            1.38,
            [1.15, 2.05, 1.5, 1.45, 1.35, 2.95, 1.68],
            rows,
            size=15,
            header_size=14,
            row_h=0.36,
            header_h=0.62,
            aligns=["c", "l", "r", "r", "c", "l", "c"],
            merges=[(frozen, 2, 6)] if frozen else None,
            fills={frozen: "FBE9E0"} if frozen else None,
            label="tier table",
        )
        textbox(
            slide,
            0.6,
            3.92,
            12.1,
            0.32,
            [{"text": self.r(s["footnote"].strip('"')), "size": 12, "color": MUTED}],
            label="tier footnote",
        )
        b = self.bullets(7)
        textbox(slide, 0.5, 4.26, 5.95, 2.62, b, size=18, bullet=True, space_after=3, label="s7 bullets")
        steps = {
            "S10": "tx_s10_promotion",
            "S15": "tx_s15_violation",
            "S16": "tx_s16_violation_freeze",
            "S18": "tx_s18_unfreeze",
        }
        hashes = {k: self.v.short(self.v.get(n)) for k, n in steps.items() if self.v.raw(n)}
        png = ASSETS / "tier-timeline.png"
        tier_chart(png, self.v.get("wallet_tier_timeline"), hashes)
        picture(slide, png, 6.55, 4.22, w=6.2)
        links = ", ".join(f"{k} {self.r('{{' + n + '|link}}')}" for k, n in steps.items() if self.v.raw(n))
        textbox(
            slide,
            6.55,
            6.68,
            6.2,
            0.4,
            [{"text": "Tier changes: " + links, "size": 12, "color": MUTED}],
            label="tier links",
        )
        self.speaker_notes(slide, 7)

    # 8 ------------------------------------------------------------------------------------------

    def slide8(self) -> None:
        slide = self.new_slide(8)
        b = self.bullets(8)
        x0, x1, ly = 0.9, 12.4, 2.12

        def xs(step: float) -> float:
            return x0 + (step - 5) / 15 * (x1 - x0)

        arrow(slide, [(x0 - 0.2, ly), (x1 + 0.25, ly)], MUTED, width=2.0)
        marks = [
            (5, "S5 mandate M1", OI_BLUE),
            (6.5, "S6/S7 pay A, capped B", OI_BLUE),
            (8, "S8 escrow, release", OI_ORANGE),
            (11.5, "S11/S12 Ghost to D, refused", OI_VERM),
            (13, "S13 stolen key reverts", OI_VERM),
            (15.5, "S15/S16 demote, freeze", OI_VERM),
            (18, "S18 owner reset, approval", OI_BLUE),
        ]
        for i, (st, lab, col) in enumerate(marks):
            cx = xs(st)
            box(slide, cx - 0.12, ly - 0.12, 0.24, 0.24, fill=col, line=WHITE, shape=MSO_SHAPE.OVAL)
            up = i % 2 == 0
            textbox(
                slide,
                cx - 0.95,
                ly - 0.72 if up else ly + 0.16,
                1.9,
                0.56,
                [{"text": lab, "size": 13, "align": "c", "bold": True}],
                anchor="b" if up else "t",
                margin=0.02,
                label=f"mark {lab}",
            )
        textbox(
            slide, x1 - 0.1, ly + 0.12, 0.6, 0.3, [{"text": "S20", "size": 11, "color": MUTED}], margin=0, label="s20"
        )
        textbox(slide, 0.5, 3.0, 5.95, 3.8, b, size=18, bullet=True, space_after=8, label="s8 bullets")
        rows = [[self.r(c) for c in r] for r in self.S[8]["table"]]
        # one transaction link per line: a hash never wraps in the middle
        rows = [r[:2] + [r[2].replace("), [", ")\n[")] + r[3:] if len(r) > 2 else r for r in rows]
        table(slide, 6.65, 3.05, [1.05, 2.45, 2.6], rows, size=13, row_h=0.5, header_h=0.36, label="demo table")
        foot = self.r(
            "All {{tx_headline}} headline claims: evidence/"
            + self.network
            + "/HEADLINE.md; verify with scripts/verify_evidence.py (backup slides list them)"
        )
        footer(slide, foot, size=12)
        self.speaker_notes(slide, 8)

    # 9 ------------------------------------------------------------------------------------------

    def slide9(self) -> None:
        slide = self.new_slide(9)
        b = self.bullets(9)
        rows = [[self.r(c) for c in r] for r in self.S[9]["table"]]
        table(
            slide,
            0.6,
            1.38,
            [2.5, 0.92, 0.92, 1.05, 0.95, 1.05, 0.85],
            rows,
            size=16,
            header_size=12,
            row_h=0.4,
            header_h=0.42,
            aligns=["l"] + ["c"] * 6,
            fills={1: "E8EEF9", 2: "E8EEF9", 3: "E8EEF9", 4: "E8EEF9"},
            bold_rows={len(rows) - 1},
            label="attack table",
        )
        failures = self.failed_attacks()
        results = [b[0], b[1].replace("Failed:", "Failed (goal reached):", 1), b[2]]
        if failures:
            results.insert(2, {"text": "; ".join(f["slide"] for f in failures), "size": 15, "color": MUTED})
        textbox(
            slide,
            0.5,
            4.98,
            8.4,
            1.86,
            [r for r in results if r],
            size=17,
            bullet=True,
            space_after=4,
            label="s9 results",
        )
        box(
            slide,
            9.15,
            1.38,
            3.6,
            2.35,
            paras=[{"text": "Compliance", "size": 17, "bold": True, "color": ACCENT_DARK}]
            + [
                {
                    "text": re.sub(r"^Compliance:\s*", "", t).capitalize() if t.startswith("Compliance:") else t,
                    "size": 15,
                    "bullet": True,
                }
                for t in b[3:5]
            ],
            space_after=4,
            margin=0.12,
            label="compliance",
        )
        layers = [
            "L1 invoice validation, detector",
            "L2 quarantine, turn-type tool sets",
            "L3 action selector, echo-and-verify",
            "L4 policy",
            "L5 signer",
            "L6 vault on-chain",
            "L7 bureau rules",
            "L8 owner console",
            "L9 autonomy operator",
        ]
        box(
            slide,
            9.15,
            3.85,
            3.6,
            2.95,
            paras=[{"text": "Nine defence layers", "size": 16, "bold": True, "color": ACCENT_DARK}]
            + [{"text": t, "size": 14} for t in layers],
            space_after=1,
            margin=0.12,
            label="layers",
        )
        commit = self.v.raw("atk_git_commit") or ""
        # a commit is printed only when the suite recorded a real one (a run before the first commit has none)
        made = f", commit {commit[:12]}" if re.fullmatch(r"[0-9a-f]{7,40}", commit) else self.r(", {{atk_run_at}}")
        foot = self.r(
            "Suite run {{atk_run_id}} on {{atk_network}}" + made + ". Stolen-key reverts "
            "on "
            + ("Base Sepolia" if self.network == "base_sepolia" else self.network)
            + ": grade gate {{tx_s13_revert_grade|link}}, replay {{tx_s13_revert_replay|link}}, frozen "
            "{{tx_s17_revert_frozen|link}}. Full table in the security self-assessment"
        )
        footer(slide, foot, size=12)
        s = self.S[9]
        paras = notes_paragraphs(s["notes_lines"])
        failed = self.v.raw("atk_failed")
        keep = []
        for p in paras:
            if p.startswith("Say this line only if Failed is 0:"):
                if failed == "0":
                    keep.append(p.split(":", 1)[1].strip())
            elif p.startswith("Say this line only if Failed is above 0:"):
                # what happened, from the suite's own checks; no fix is claimed that the evidence does not show
                keep += [f["notes"] for f in failures]
            else:
                keep.append(p)
        self.speaker_notes(slide, 9, keep)

    def failed_attacks(self) -> list[dict[str, str]]:
        """Each Failed attack described from ``attack_suite.json``: what it tried, which checks held, which not."""
        try:
            suite = self.v.ctx.load("evidence/attack_suite.json")
        except Exception:  # noqa: BLE001 - no suite file: the placeholders report it
            return []
        out = []
        for r in suite.get("results") or []:
            if r.get("actual") != "Failed":
                continue
            held = [c for c in r.get("checks") or [] if c.get("ok")]
            broke = [c for c in r.get("checks") or [] if not c.get("ok")]
            lost = rd.fmt_value(r.get("fundsLost") or "0", "usdc")

            def say(c: dict[str, Any]) -> str:
                d = c.get("detail")
                if isinstance(d, list):
                    d = ", ".join(str(x) for x in d) if d else "none"
                return f"{c['check']} ({d})" if d not in (None, "", True, False) else c["check"]

            what = f"{r['id']}, {r.get('title', '').rstrip('.').lower()} (expected {r.get('expected')})"
            held_t = "; ".join(say(c) for c in held)
            broke_t = "; ".join(say(c) for c in broke)
            seen = broke[0].get("detail") if broke else None
            seen_n = len(seen) if isinstance(seen, list) else None
            short = (
                f"{r['id']} ({r.get('class', '').replace('_', ' ')}): {lost} USDC lost"
                + (f", {held_t}" if held_t else "")
                + (
                    f'; counted Failed: check "{broke[0]["check"]}" '
                    + (
                        f"saw {seen_n} {plural(seen_n, 'transaction', 'transactions')} ({', '.join(map(str, seen))})"
                        if seen_n
                        else "did not hold"
                    )
                    if broke
                    else ""
                )
            )
            notes_t = f"Counted Failed: {what}. What held: {held_t or 'none'}; {lost} USDC was lost. " + (
                f'The check "{broke[0]["check"]}" saw agent-key transactions '
                f"({', '.join(map(str, seen))}), so the suite counts it as reaching its goal; the check "
                "cannot tell whether the attack caused those transactions."
                if seen_n
                else f"What did not hold: {broke_t}, so the suite counts it as reaching its goal."
            )
            out.append({"slide": short, "notes": notes_t})
        return out

    # 10 -----------------------------------------------------------------------------------------

    def s5_dialogue(self) -> dict[str, Any] | None:
        try:
            return self.v.ctx.load("evidence/<network>/dialogue/S5.json")
        except Exception:  # noqa: BLE001 - older evidence sets have no dialogue file
            return None

    def slide10(self) -> None:
        slide = self.new_slide(10)
        b = self.bullets(10)
        mode = self.v.raw("llm_s5_mode") or ""
        mock = mode.lower().startswith("mock")
        d = self.s5_dialogue()
        box(slide, 0.6, 1.38, 7.6, 4.14, fill="F7F8FC", line=CARD_LINE, radius=0.03)
        label = f"S5 chat: owner lines scripted; model {mode}" + (", a scripted mock, not a live LLM" if mock else "")
        textbox(slide, 0.75, 1.42, 7.3, 0.34, [{"text": label, "size": 13, "color": MUTED}], label="chat label")

        def usdc(v: Any) -> str:
            return rd.fmt_value(v, "usdc")

        def owner(y: float, h: float, text: str, size: float = 15) -> None:
            box(
                slide,
                2.55,
                y,
                5.5,
                h,
                fill=OI_BLUE,
                line=None,
                radius=0.25,
                paras=[{"text": text, "size": size, "color": WHITE}],
                anchor="m",
                margin=0.14,
                label="owner",
            )

        def agent(y: float, h: float, head: str, text: str, line: str = CARD_LINE, size: float = 14) -> None:
            box(
                slide,
                0.75,
                y,
                6.1,
                h,
                fill=WHITE,
                line=line,
                line_w=1.5,
                radius=0.15,
                paras=[{"text": head, "size": 13, "bold": True, "color": ACCENT_DARK}, {"text": text, "size": size}],
                space_after=1,
                anchor="m",
                margin=0.14,
                label=f"agent {head}",
            )

        turns = [t for t in (d or {}).get("turns") or [] if "owner" in t]  # the last entry is the confirmation
        cards = [(t, c) for t in turns for c in t.get("cards") or []]
        clar = next((c for _, c in cards if c.get("type") == "clarification"), None)
        mandates = [c for _, c in cards if c.get("type") == "mandate" and c.get("values")]
        facts: list[list[str]] = []
        if len(turns) >= 3 and clar and len(mandates) >= 2:
            # the real S5 exchange of this run, shortened to names, grades and payout wallets
            cands = clar.get("candidates") or []
            listed = "; ".join(f"{c['name']} ({c['grade']}, {self.v.short(c['liveWallet'])})" for c in cands)
            first, last = mandates[0]["values"], mandates[-1]["values"]
            owner(1.78, 0.62, turns[0]["owner"])
            lead = next(
                (t["reply"].split(":", 1)[0] for t in turns if clar in (t.get("cards") or [])),
                f"{len(cands)} providers match",
            )
            agent(2.48, 1.2, "Agent: forced clarification (shortened)", f"{lead}: {listed}. Which should I use?")
            owner(3.76, 0.6, turns[1]["owner"], size=14)
            owner(4.44, 0.4, turns[-1]["owner"], size=14)
            diff = (
                f"total {usdc(first['maxTotal'])} -> {usdc(last['maxTotal'])} USDC"
                if first["maxTotal"] != last["maxTotal"]
                else f"total {usdc(last['maxTotal'])} USDC"
            )
            agent(
                4.92,
                0.56,
                "Agent: mandate card shows the diff",
                f"{diff}; {len(last['payeeAgentIds'])} payees, {usdc(last['maxPerTx'])} per payment",
                line=OI_ORANGE,
            )
            facts = [
                ["Owner turns", str(len(turns))],
                ["Model", (d or {}).get("llm", {}).get("model") or mode],
                ["Forced clarifications", str(sum(1 for _, c in cards if c.get("type") == "clarification"))],
                [
                    "Rule-based fallback used",
                    "yes" if any((t.get("llm") or {}).get("fallback") for t in turns) else "no",
                ],
                ["Mandate payees", str(len(last["payeeAgentIds"]))],
                ["Per payment / total", f"{usdc(last['maxPerTx'])} / {usdc(last['maxTotal'])} USDC"],
            ]
        else:
            owner(1.82, 0.95, b[0].split(":", 1)[1].strip().strip('"'), size=17)
            agent(2.92, 0.95, "Agent (forced clarification)", b[1].split(":", 1)[1].strip(), size=16)
            owner(4.0, 0.5, b[2].split("->")[0].split(":", 1)[1].strip().strip('"'), size=17)
            agent(4.6, 0.62, "Agent", "Mandate card shows the diff", line=OI_ORANGE, size=16)
        textbox(
            slide,
            0.5,
            5.62,
            7.8,
            1.25,
            [x for x in b[3:] if x],
            size=18,
            bullet=True,
            space_after=4,
            label="s10 bullets",
        )
        ev = [
            ("Requests evaluated", "{{de_cases}}"),
            ("Correct mandates", "{{de_correct_mandate_rate}}"),
            ("Clarification precision", "{{de_clarification_precision}}"),
            ("Clarification recall", "{{de_clarification_recall}}"),
            ("Correction success", "{{de_correction_success_rate}}"),
            ("Chinese requests correct", "{{de_chinese_correct}} of {{de_chinese_cases}}"),
            ("Hostile requests safe", "{{de_hostile_safe_rate}}"),
            ("Median turns", "{{de_turns_median}}"),
        ]
        evaluated = [[a, self.r(z)] for a, z in ev]
        banner_y = 5.0
        if all(MISSING not in z for _, z in evaluated):
            banner_y = 1.38 + 0.42 * (1 + len(evaluated)) + 0.18  # below the table, never over its last row
            table(
                slide,
                8.5,
                1.38,
                [2.85, 1.4],
                [["Dialogue evaluation", "Result"]] + evaluated,
                size=15,
                row_h=0.42,
                header_h=0.42,
                aligns=["l", "r"],
                label="dialogue eval",
            )
        elif facts:
            # no dialogue evaluation in this evidence set: the S5 exchange itself, from dialogue/S5.json
            banner_y = 1.38 + 0.42 + 0.5 * len(facts) + 0.18
            table(
                slide,
                8.5,
                1.38,
                [2.45, 1.8],
                [["S5 in this run", "Value"]] + facts,
                size=15,
                row_h=0.5,
                header_h=0.42,
                aligns=["l", "r"],
                label="s5 facts",
            )
        box(
            slide,
            8.5,
            banner_y,
            4.25,
            0.62,
            fill="FFF4DD",
            line=OI_ORANGE,
            radius=0.2,
            paras=[{"text": "Banner when the model is down: AI offline: rule-based mode", "size": 14}],
            anchor="m",
            margin=0.12,
            label="banner",
        )
        paras = notes_paragraphs(self.S[10]["notes_lines"])
        if mock:
            paras = [
                p.replace(
                    "In the run the owner's lines are scripted and the model's side is live.",
                    f"In this dry run both sides are scripted: the owner's lines, and a mock model "
                    f"({mode}) instead of a live LLM.",
                )
                for p in paras
            ]
        self.speaker_notes(slide, 10, paras)

    # 11 -----------------------------------------------------------------------------------------

    def bound_values(self, b: list[str]) -> tuple[str, str]:
        """Stolen-key bound and grade-A exposure limit, from config; deck.md's wording is checked against them."""
        scen = yaml.safe_load((REPO / "config" / "scenario.yaml").read_text(encoding="utf-8"))
        pol = yaml.safe_load((REPO / "config" / "policy.yaml").read_text(encoding="utf-8"))
        # spec 2.5: min(remaining mandate budget, window cap, free balance); the mandate total binds in the demo
        stolen = rd.fmt_value(scen["mandate"]["max_total"], "usdc")
        exposure = rd.fmt_value(pol["exposure_limit"]["A"], "usdc")
        said = re.findall(r"(\d+\.\d\d) USDC", " ".join(b[:2]))
        if said[:2] != [stolen, exposure]:
            WARNINGS.append(f"slide 11: deck.md says {said[:2]}, config gives {[stolen, exposure]}")
        return stolen, exposure

    def slide11(self) -> None:
        slide = self.new_slide(11)
        b = self.bullets(11)
        stolen, exposure = self.bound_values(self.S[11]["bullets"])
        tiles = [
            (f"{stolen} USDC", "stolen-key bound, only to approved payees"),
            (f"{exposure} USDC", "single-name exposure limit, grade A"),
        ]
        if self.have("benign_false_positives", "benign_tasks"):
            tiles.append((self.r("{{benign_false_positives}} / {{benign_tasks}}"), "false positives on normal tasks"))
        n = len(tiles)
        w = (12.15 - (n - 1) * 0.2) / n
        for i, (big, small) in enumerate(tiles):
            box(
                slide,
                0.6 + i * (w + 0.2),
                1.4,
                w,
                1.45,
                paras=[
                    {"text": big, "size": 34, "bold": True, "color": OI_BLUE, "align": "c"},
                    {"text": small, "size": 16, "align": "c"},
                ],
                space_after=4,
                anchor="m",
                margin=0.12,
                label=f"s11 tile {i}",
            )
        extra: list[str] = []
        try:
            rounds = self.v.ctx.load("evidence/backtest.json")["rounds"]
            pcts = [float(rd.fmt_value(r["falseAdverseBps"], "bps_pct").rstrip("%")) for r in rounds]
        except Exception:  # noqa: BLE001 - no backtest in this evidence set
            pcts = []
        if pcts:
            cd = CategoryChartData()
            cd.categories = [str(i + 1) for i in range(len(pcts))]
            cd.add_series("Honest agents graded C or D", pcts)
            gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, In(0.6), In(3.02), In(5.6), In(3.78), cd)
            ch = gf.chart
            ch.has_legend = False
            ch.has_title = True
            ch.chart_title.text_frame.text = "Honest agents graded C or D (backtest)"
            tr = ch.chart_title.text_frame.paragraphs[0].runs[0]
            tr.font.size, tr.font.bold, tr.font.name = Pt(16), True, FONT
            tr.font.color.rgb = rgb(INK)
            plot = ch.plots[0]
            plot.gap_width = 60
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.number_format, dl.number_format_is_linked = '0.0"%"', False
            dl.position = XL_LABEL_POSITION.OUTSIDE_END
            dl.font.size, dl.font.name = Pt(14), FONT
            ser = plot.series[0]
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = rgb(OI_BLUE)
            va = ch.value_axis
            va.maximum_scale = max(1.0, max(pcts) * 1.25)
            va.minimum_scale = 0
            va.has_major_gridlines = False
            va.visible = False
            ca = ch.category_axis
            ca.tick_labels.font.size, ca.tick_labels.font.name = Pt(14), FONT
            ca.format.line.color.rgb = rgb(CARD_LINE)
            ca.has_title = True
            ca.axis_title.text_frame.text = "Round (early rounds: bootstrap, partners unrated)"
            at = ca.axis_title.text_frame.paragraphs[0].runs[0]
            at.font.size, at.font.bold, at.font.name = Pt(14), False, FONT
            at.font.color.rgb = rgb(MUTED)
            extra.append(
                self.r(
                    "The early bars are the bootstrap: every agent starts at NEW, so honest agents start at C or "
                    "D ({{bt_false_adverse_round1_pct}} in round 1); by the last round it is {{bt_false_adverse_pct}}."
                )
            )
        else:
            box(slide, 0.6, 3.15, 5.6, 3.6, paras=["No backtest in this evidence set"], label="no chart")
        stats = [
            {"text": self.r("Backtest on {{bt_agents}} synthetic agents"), "size": 18, "bold": True},
            {"text": self.r("AUC {{bt_auc}}, KS {{bt_ks}}"), "size": 18, "space_after": 2},
            {
                "text": "A consistency check on our own simulator, not an out-of-sample validation",
                "size": 15,
                "color": MUTED,
            },
            {"text": self.r("Largest honest score change from smears: {{bt_smear_max_abs_delta}} points"), "size": 18},
            {"text": b[4], "size": 18, "bold": True, "color": ACCENT_DARK},
        ]
        stats = [s for s in stats if s["text"] and MISSING not in s["text"]]
        textbox(slide, 6.5, 3.15, 6.25, 3.6, stats, space_after=10, label="s11 stats")
        footer(
            slide,
            self.r(
                "Cost of the run: {{eth_spent_total}} ETH of testnet fees; a payment uses about "
                "{{gas_median_pay}} gas. Bounds: docs/design/spec.md 2.5-2.6, config/policy.yaml"
            ),
            size=12,
        )
        extra.append(
            "Measured on our own simulator, AUC and KS are a consistency check, not an out-of-sample validation."
        )
        self.speaker_notes(slide, 11, extra=extra)

    # 12 -----------------------------------------------------------------------------------------

    def slide12(self) -> None:
        slide = self.new_slide(12)
        b = self.bullets(12)
        not_new = [x.split(":", 1)[1].strip() for x in b if x.startswith("Not new")]
        new = [x.split(":", 1)[1].strip() for x in b if x.startswith("New")]
        for i, (head, items, col) in enumerate((("Not new", not_new, MUTED), ("New", new, OI_BLUE))):
            box(
                slide,
                0.6 + i * 6.15,
                1.38,
                5.95,
                1.62,
                paras=[{"text": head, "size": 18, "bold": True, "color": col}]
                + [{"text": t[0].upper() + t[1:], "size": 16, "bullet": True} for t in items],
                space_after=3,
                margin=0.12,
                label=f"s12 {head}",
            )
        rows = [[self.r(c) for c in r] for r in self.S[12]["table"]]
        widths = [3.75, 4.1, 4.28]
        table(slide, 0.6, 3.12, widths, rows, size=13, row_h=0.46, header_h=0.36, label="related")
        bottom = 3.12 + table_height(widths, rows, 13, 0.46, 0.36)
        if bottom > 6.15:  # the estimate is conservative (rows are rarely as tall)
            WARNINGS.append(f"slide 12: related-work table reaches {bottom:.2f} in, the claim line starts at 6.2")
        textbox(
            slide,
            0.5,
            6.2,
            12.3,
            0.66,
            [
                {
                    "text": "Our narrow claim: shared on-chain behaviour grades for "
                    "agents, from outcomes anyone can verify, that change both whom a wallet pays and how much "
                    "freedom the wallet itself has.",
                    "size": 15,
                    "bold": True,
                    "color": ACCENT_DARK,
                }
            ],
            label="claim",
        )
        self.speaker_notes(slide, 12)

    # 13 -----------------------------------------------------------------------------------------

    def slide13(self) -> None:
        slide = self.new_slide(13)
        b = [x for x in self.bullets(13) if self.blockscout_verified or "Blockscout" not in x]
        step = min(1.1, 5.3 / max(len(b), 1))
        for i, t in enumerate(b):
            y = 1.45 + i * step
            box(
                slide,
                0.6,
                y + 0.12,
                0.6,
                0.6,
                fill=ACCENT,
                line=None,
                shape=MSO_SHAPE.OVAL,
                paras=[{"text": str(i + 1), "size": 18, "bold": True, "color": WHITE, "align": "c"}],
                anchor="m",
                margin=0,
                label="num",
            )
            box(slide, 1.35, y, 6.4, step - 0.2, paras=[t], size=18, anchor="m", margin=0.12, label=f"s13 card {i}")
        match, total = self.v.get("grade_changes_recompute_match"), self.v.get("grade_changes_total")
        box(slide, 8.1, 1.4, 4.65, 3.95, fill=CARD, line=CARD_LINE, radius=0.03)
        textbox(
            slide,
            8.25,
            1.5,
            4.4,
            0.45,
            [{"text": "Permissionless recompute", "size": 18, "bold": True, "color": ACCENT_DARK}],
            label="rc head",
        )
        textbox(
            slide,
            8.25,
            2.05,
            4.4,
            1.1,
            [{"text": f"{match} / {total}", "size": 54, "bold": True, "color": OI_BLUE, "align": "c"}],
            anchor="m",
            label="rc big",
        )
        textbox(
            slide,
            8.25,
            3.15,
            4.4,
            0.7,
            [
                {
                    "text": "grade changes recompute to MATCH from public logs and the published evidence files",
                    "size": 15,
                    "align": "c",
                }
            ],
            label="rc sub",
        )
        box(
            slide,
            8.3,
            4.15,
            4.25,
            0.95,
            fill="1E2235",
            line=None,
            radius=0.08,
            paras=[
                {"text": "$ acb bureau recompute --tx <hash>", "font": MONO, "size": 13, "color": "E6E8F2"},
                {"text": "MATCH", "font": MONO, "size": 13, "bold": True, "color": OI_GREEN},
            ],
            anchor="m",
            margin=0.12,
            space_after=2,
            label="rc cmd",
        )
        if self.network == "base_sepolia":
            strip = self.r(
                "ERC-8004 version {{erc8004_version}} checked at block {{erc8004_checked_block}}; "
                "AgentCreditBureau {{addr_acb|blockscout}}; wallet vault {{addr_vault_wallet|blockscout}}; "
                "upgrade delay {{acb_upgrade_delay_s}} s"
            )
        else:
            strip = self.r(
                f"On {self.network} the ERC-8004 registries are local Lite copies (version check is "
                "a Base Sepolia step); AgentCreditBureau {{addr_acb|short}}; wallet vault "
                "{{addr_vault_wallet|short}}; upgrade delay {{acb_upgrade_delay_s}} s"
            )
        footer(slide, strip, size=12)
        s = self.S[13]
        paras = notes_paragraphs(s["notes_lines"])
        if not self.blockscout_verified:
            paras = [p.replace(" Contract sources are verified on Blockscout.", "") for p in paras]
        self.speaker_notes(slide, 13, paras)

    # 14 -----------------------------------------------------------------------------------------

    def slide14(self) -> None:
        slide = self.new_slide(14)
        b = self.bullets(14)
        limits = [x.split(":", 1)[1].strip() for x in b if x.startswith("Limits")]
        nxt = [t.strip() for x in b if x.startswith("Next") for t in x.split(":", 1)[1].split(",")]
        for i, (head, items, col) in enumerate((("Known limits", limits, OI_VERM), ("Next", nxt, OI_BLUE))):
            x = 0.6 + i * 6.15
            box(slide, x, 1.45, 5.95, 4.1, fill=CARD, line=CARD_LINE, radius=0.03)
            box(
                slide,
                x + 0.25,
                1.65,
                0.5,
                0.5,
                fill=col,
                line=None,
                shape=MSO_SHAPE.OVAL,
                paras=[{"text": "!" if i == 0 else ">", "size": 18, "bold": True, "color": WHITE, "align": "c"}],
                anchor="m",
                margin=0,
                check=False,
            )
            textbox(
                slide,
                x + 0.9,
                1.62,
                4.8,
                0.56,
                [{"text": head, "size": 22, "bold": True, "color": INK}],
                anchor="m",
                label=f"s14 {head}",
            )
            textbox(
                slide,
                x + 0.25,
                2.4,
                5.5,
                3.05,
                [{"text": t[0].upper() + t[1:], "size": 22, "bullet": True} for t in items],
                space_after=14,
                label=f"s14 items {head}",
            )
        box(
            slide,
            0.6,
            5.8,
            12.1,
            0.62,
            fill=WHITE,
            line=CARD_LINE,
            radius=0.2,
            paras=[{"text": self.timings(), "size": 16, "color": MUTED, "align": "c"}],
            anchor="m",
            label="timings",
        )
        self.speaker_notes(slide, 14)

    def timings(self) -> str:
        """deck.md's timing line, its numbers checked against config/scorecard.yaml and config/networks.yaml."""
        m = re.search(r'Small grey line at the bottom:\s*"([^"]+)"', self.S[14]["visual"])
        if not m:
            raise SystemExit("deck.md slide 14: the timing line is missing from the Visual text")
        line = m.group(1)
        sc = yaml.safe_load((REPO / "config" / "scorecard.yaml").read_text(encoding="utf-8"))["profiles"]
        nets = yaml.safe_load((REPO / "config" / "networks.yaml").read_text(encoding="utf-8"))
        delays = {
            str(v.get("upgrade_delay_s"))
            for v in (nets.get("networks") or nets).values()
            if isinstance(v, dict) and "upgrade_delay_s" in v
        }
        prod_days = sc["production"]["a_min_age_s"] / 86400
        want = [f"age gate {sc['demo']['a_min_age_s']} s", f"(production {prod_days:g} days)"]
        want += [f"grade upgrade delay {d} s" for d in delays]
        for w in want:
            if w not in line:
                WARNINGS.append(f"slide 14: deck.md timing line does not say {w!r} (config)")
        return line

    # 15 -----------------------------------------------------------------------------------------

    def slide15(self) -> None:
        slide = self.prs.slides[3]
        b = self.bullets(15)
        box(slide, 6.75, 5.05, 6.2, 2.1, fill="0B0E26", line=None, radius=0.08, alpha=78)
        if qr_png(ASSETS / "qr-repo.png", REPO_URL):
            picture(slide, ASSETS / "qr-repo.png", 6.9, 5.2, w=1.8, h=1.8)
        url = b[1].split("/", 1)
        textbox(
            slide,
            8.83,
            5.2,
            4.07,
            1.85,
            [
                {
                    "text": f"[{url[0]}/{url[1]}]({REPO_URL})" if len(url) == 2 else b[1],
                    "size": 14,
                    "bold": True,
                    "color": WHITE,
                },
                {"text": b[2], "size": 14, "color": "D7DBEE"},
            ],
            space_after=6,
            label="thanks text",
        )
        self.draft_tag(slide, dark=True)
        self.speaker_notes(slide, 15)

    # backup ---------------------------------------------------------------------------------------

    def backup(self) -> None:
        text = self.v.raw("headline_table")
        if text is None:
            return
        lines = [ln for ln in text.splitlines() if ln.startswith("|")]
        head = [c.strip() for c in lines[0].strip("|").split("|")]
        body = [[c.strip() for c in ln.strip("|").split("|")] for ln in lines[2:]]
        want = ["#", "Claim", "Step", "Function", "Reasons", "Transaction"]
        idx = [head.index(c) for c in want]
        rows = []
        for r in body:
            cells = [r[i] for i in idx]
            m = re.search(r"0x[0-9a-f]{64}", cells[-1])
            cells[-1] = self.v.link_md(m.group(0)) if m else cells[-1]
            cells[3] = cells[3].strip("`")
            rows.append(cells)
        half = (len(rows) + 1) // 2
        for part, chunk in enumerate((rows[:half], rows[half:]), start=1):
            if not chunk:
                continue
            slide = self.prs.slides.add_slide(self.layout)
            title(slide, f"Backup {part}/2: headline claims and their transactions")
            self.draft_tag(slide)
            table(
                slide,
                0.5,
                1.35,
                [0.45, 6.15, 0.7, 1.75, 1.55, 1.75],
                [want] + chunk,
                size=12,
                row_h=0.36,
                header_h=0.34,
                aligns=["c", "l", "c", "l", "l", "l"],
                label=f"backup {part}",
            )
            footer(slide, f"From evidence/{self.network}/HEADLINE.md (not counted in the 10 minutes; for questions)")
            notes(slide, "Backup slide for questions; not part of the timed talk.")

    def build(self) -> None:
        self.slide1()
        self.slide2()
        self.slide3()
        for n in range(4, 15):
            getattr(self, f"slide{n}")()
        self.slide15()
        self.backup()
        lst = self.prs.slides._sldIdLst
        ids = list(lst)
        thanks = ids[3]
        lst.remove(thanks)
        lst.insert(14, thanks)


# ---------------------------------------------------------------------------------------------------
# main
# ---------------------------------------------------------------------------------------------------


LEAK_RE = re.compile(MISSING + r"|\bn/a\b|uncommitted|\[Presenter|\{\{")


def leftovers(prs) -> list[tuple[str, str]]:
    """Missing values, unfinished prompts and unrendered placeholders anywhere on a slide or in its notes."""
    found = []

    def texts(shapes):
        for sh in shapes:
            if sh.has_text_frame:
                yield sh.text_frame.text
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    for cell in row.cells:
                        yield cell.text_frame.text
            if sh.shape_type == 6:  # group
                yield from texts(sh.shapes)

    for i, slide in enumerate(prs.slides, start=1):
        items = [("slide", t) for t in texts(slide.shapes)]
        if slide.has_notes_slide:
            items.append(("notes", slide.notes_slide.notes_text_frame.text))
        for kind, t in items:
            for m in LEAK_RE.finditer(t):
                found.append((f"slide {i} {kind}", t[max(0, m.start() - 40) : m.end() + 40].replace("\n", " ")))
    return found


def default_evidence() -> Path:
    if (REPO / "evidence" / "base_sepolia" / "evidence.json").is_file():
        return REPO / "evidence"
    print(
        "note: evidence/base_sepolia/evidence.json not found; using the local Anvil set in "
        f"{STALE.relative_to(REPO).as_posix()}",
        file=sys.stderr,
    )
    return STALE


def detect_network(root: Path) -> str:
    if (root / "evidence.json").is_file():
        return json.loads((root / "evidence.json").read_text(encoding="utf-8-sig"))["network"]
    for net in ("base_sepolia", "anvil"):
        if (root / net / "evidence.json").is_file():
            return net
    raise SystemExit(f"no <network>/evidence.json under {root}")


def soffice() -> str | None:
    for c in (
        shutil.which("soffice"),
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        "/usr/bin/soffice",
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    ):
        if c and Path(c).is_file():
            return c
    return None


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--evidence", type=Path, default=None, help="evidence folder (default: see module docstring)")
    ap.add_argument("--network", default=None, help="network name (default: detected from the folder)")
    ap.add_argument("--draft", action="store_true", help="allow missing values on Base Sepolia evidence")
    ap.add_argument(
        "--blockscout-verified",
        action="store_true",
        help="keep the 'verified on Blockscout' bullet of slide 13 (only after forge verify-contract)",
    )
    ap.add_argument("--out", type=Path, default=OUT)
    ap.add_argument("--pdf", action="store_true", help="also write a PDF copy with LibreOffice (soffice)")
    ap.add_argument(
        "--previews", action="store_true", help="also render slide PNGs to deck/previews/ (needs --pdf and pdftoppm)"
    )
    args = ap.parse_args(argv)

    root = (args.evidence or default_evidence()).resolve()
    network = args.network or detect_network(root)
    draft = args.draft or network != "base_sepolia"
    entries, groups = rd.load_registry()
    ctx = EvidenceContext(network, groups, root)
    vals = Values(ctx, entries)
    slides = parse_deck(SOURCE.read_text(encoding="utf-8-sig"))
    if not TEMPLATE.is_file():
        raise SystemExit(f"template not found: {TEMPLATE}")
    ASSETS.mkdir(parents=True, exist_ok=True)
    prs = Presentation(str(TEMPLATE))
    Builder(prs, slides, vals, network, draft, args.blockscout_verified).build()
    # the template's document properties carry its default Chinese title and its last editor's name
    cp = prs.core_properties
    cp.title = "Agent Credit Bureau - Team Mongolia"
    cp.subject = "2026 FinTechathon, International AI Track, Topic F: Agent Wallet"
    cp.author = "Team Mongolia (National University of Mongolia)"
    cp.last_modified_by = "Team Mongolia"
    cp.revision = 1
    cp.created = cp.modified = datetime.now(UTC).replace(tzinfo=None, microsecond=0)

    if vals.missing:
        print(
            f"{len(vals.missing)} value(s) not in this evidence set (their tiles, rows and sentences are left out):",
            file=sys.stderr,
        )
        for k, why in sorted(vals.missing.items()):
            print(f"  {k}: {why}", file=sys.stderr)
        required = [k for k in vals.missing if not k.startswith(OPTIONAL_PREFIXES)]
        if required and not draft:
            print(
                f"nothing written: Base Sepolia evidence must resolve {', '.join(sorted(required))} (or pass --draft)",
                file=sys.stderr,
            )
            return 1
    leaks = leftovers(prs)
    if leaks:
        print("nothing written: text that must not reach a slide:", file=sys.stderr)
        for where, what in leaks:
            print(f"  {where}: {what}", file=sys.stderr)
        return 1
    for w in WARNINGS:
        print(f"layout warning: {w}", file=sys.stderr)

    args.out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(args.out))
    print(
        f"wrote {args.out.relative_to(REPO).as_posix() if args.out.is_relative_to(REPO) else args.out} "
        f"({len(prs.slides)} slides, evidence {root.name}, network {network}{', DRAFT' if draft else ''})"
    )

    if args.pdf or args.previews:
        exe = soffice()
        if not exe:
            print("LibreOffice (soffice) not found: no PDF or previews", file=sys.stderr)
            return 0
        subprocess.run(
            [exe, "--headless", "--convert-to", "pdf", "--outdir", str(args.out.parent), str(args.out)],
            check=True,
            capture_output=True,
            timeout=300,
        )
        pdf = args.out.with_suffix(".pdf")
        print(f"wrote {pdf.name}")
        if args.previews:
            tool = shutil.which("pdftoppm")
            if not tool:
                print("pdftoppm not found: no previews", file=sys.stderr)
                return 0
            PREVIEWS.mkdir(parents=True, exist_ok=True)
            for old in PREVIEWS.glob("slide-*.png"):
                old.unlink()
            subprocess.run([tool, "-png", "-r", "80", str(pdf), str(PREVIEWS / "slide")], check=True, timeout=300)
            print(f"wrote {len(list(PREVIEWS.glob('slide-*.png')))} previews to deck/previews/")
    return 0


if __name__ == "__main__":
    sys.exit(main())
