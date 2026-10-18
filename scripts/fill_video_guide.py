"""Fill the organiser's "Demo Video | Production & Submission Guide" form from the rendered video script.

The entries come from section 0 of ``docs/deliverables/out/video-script.md`` (render it first with
``scripts/render_deliverables.py``), so the form says what the rendered video shows. The organiser's
template is read from ``private/templates/demo-video-guide-en.docx`` (not redistributed in this
repository); the filled form is written to ``docs/deliverables/demo-video-guide/`` as DOCX and, when
LibreOffice is installed, as PDF.

    uv run --with python-docx python scripts/fill_video_guide.py
"""

from __future__ import annotations

import re
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "private" / "templates" / "demo-video-guide-en.docx"
SCRIPT = ROOT / "docs" / "deliverables" / "out" / "video-script.md"
OUT = ROOT / "docs" / "deliverables" / "demo-video-guide" / "InternationalAI-TeamMongolia-DemoVideoGuide.docx"

# the template's "Write here" lines, in order, and the guide row each one answers
ROWS = [
    "Project and team introduction",
    "End-to-end demonstration of the core scenario",
    "Safety mechanisms, innovation and results",
    "Use 16:9 landscape and preferably 1080p",
    "Make interface text, narration and English subtitles clearly legible",
    "No copyrighted or sensitive material",
    "Suggested filename: Track-TeamName-DemoVideo",
]
HEADER = {"队伍名称 / Team": "Team", "参赛赛道 / Track": "Track", "提交日期 / Date": "Date"}


def plain(cell: str) -> str:
    """Markdown cell text without code ticks."""
    return re.sub(r"`([^`]*)`", r"\1", cell).strip()


def section0(text: str) -> dict[str, str]:
    """The two-column rows of section 0 of the rendered video script, keyed by their first cell."""
    part = text.split("## 0.", 1)[1].split("\n## 1.", 1)[0]
    rows: dict[str, str] = {}
    for line in part.splitlines():
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if len(cells) == 2 and not set(cells[0]) <= {"-"} and cells[0] not in ("Field", "Guide row"):
            rows[cells[0]] = plain(cells[1])
    return rows


def main() -> int:
    import docx
    from docx.shared import RGBColor

    if not TEMPLATE.is_file():
        print(f"template not found: {TEMPLATE}", file=sys.stderr)
        return 1
    text = SCRIPT.read_text(encoding="utf-8")
    if "{{" in text:
        print(f"{SCRIPT}: unrendered placeholder; run scripts/render_deliverables.py first", file=sys.stderr)
        return 1
    rows = section0(text)
    missing = [r for r in [*ROWS, *HEADER.values()] if r not in rows]
    if missing:
        print(f"section 0 of the video script lacks: {missing}", file=sys.stderr)
        return 1
    doc = docx.Document(str(TEMPLATE))
    blanks = [p for p in doc.paragraphs if p.text.strip() == "Write here"]
    if len(blanks) != len(ROWS):
        print(f"the template has {len(blanks)} 'Write here' lines, expected {len(ROWS)}", file=sys.stderr)
        return 1
    ink = RGBColor(0x1E, 0x22, 0x35)  # the template prints "Write here" in placeholder grey; answers are dark
    for para, row in zip(blanks, ROWS, strict=True):
        runs = para.runs
        runs[0].text = rows[row]
        runs[0].font.color.rgb = ink
        for r in runs[1:]:
            r.text = ""
    for table in doc.tables:
        for tr in table.rows:
            label = tr.cells[0].text.strip()
            if label in HEADER:
                cell = tr.cells[1]
                runs = cell.paragraphs[0].runs
                runs[0].text = rows[HEADER[label]]
                runs[0].font.color.rgb = ink
                for r in runs[1:]:
                    r.text = ""
    cp = doc.core_properties
    cp.title = "Demo Video Guide - Agent Credit Bureau - Team Mongolia"
    cp.author = cp.last_modified_by = "Team Mongolia"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"wrote {OUT.relative_to(ROOT)}")
    soffice = shutil.which("soffice") or shutil.which("soffice.exe")
    for guess in (r"C:\Program Files\LibreOffice\program\soffice.exe",):
        if soffice is None and Path(guess).is_file():
            soffice = guess
    if soffice:
        subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(OUT.parent), str(OUT)],
            check=True,
            capture_output=True,
        )
        print(f"wrote {OUT.with_suffix('.pdf').relative_to(ROOT)}")
    else:
        print("LibreOffice (soffice) not found: no PDF copy", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
